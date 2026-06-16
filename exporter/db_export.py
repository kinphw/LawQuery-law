"""
기존 법 DB(ldb_<코드>) → data dict (reader와 동일 형태) / 엑셀.
GUI '기존 법 불러오기'와 CLI `export` 가 공유하는 읽기 엔진.
"""
import pymysql

from common.db import get_connection
from common.schema_map import SHEETS, LOAD_ORDER


def code_of(dbname: str) -> str:
    return dbname[4:] if dbname.startswith("ldb_") else dbname


def list_law_dbs(target: str = "dev") -> list[str]:
    """db_a 테이블을 가진 ldb_* DB만(법령 스키마). ldb_auth·ldb_i(해석) 제외."""
    conn = get_connection(target=target)
    laws: list[str] = []
    try:
        with conn.cursor() as cur:
            cur.execute("SHOW DATABASES LIKE %s", ("ldb\\_%",))
            for (db,) in cur.fetchall():
                cur.execute(
                    "SELECT COUNT(*) FROM information_schema.tables "
                    "WHERE table_schema=%s AND table_name='db_a'",
                    (db,),
                )
                if cur.fetchone()[0]:
                    laws.append(db)
    finally:
        conn.close()
    return laws


def _order(cols: list[str]) -> str:
    if "seq" in cols:
        return " ORDER BY `seq`"
    if "id" in cols:
        return " ORDER BY `id`"
    return ""


def read_law(code: str, target: str = "dev") -> dict:
    """ldb_<code> → {sheet: [row dict, ...]} (적재 컬럼만, 적재 순서 유지)."""
    conn = get_connection(database=f"ldb_{code}", target=target)
    data: dict = {}
    try:
        with conn.cursor(pymysql.cursors.DictCursor) as cur:
            for sheet, (table, cols) in SHEETS.items():
                collist = ", ".join(f"`{c}`" for c in cols)
                try:
                    cur.execute(f"SELECT {collist} FROM `{table}`" + _order(cols))
                    data[sheet] = [dict(r) for r in cur.fetchall()]
                except pymysql.err.ProgrammingError:
                    data[sheet] = []  # 테이블 없음(penalty 미사용 등)
    finally:
        conn.close()
    return data


def write_workbook(data: dict, out_path: str) -> str:
    """data dict → 템플릿 형식 엑셀. (DB export·GUI 내보내기 공용)"""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for sheet in LOAD_ORDER:
        _table, cols = SHEETS[sheet]
        ws = wb.create_sheet(title=sheet)
        ws.append(cols)
        for row in data.get(sheet, []):
            ws.append([row.get(c) for c in cols])
        ws.freeze_panes = "A2"
    wb.save(out_path)
    return out_path


def export_to_excel(code: str, out_path: str, target: str = "dev"):
    """기존 법을 템플릿과 동일한 형식의 엑셀로 내보낸다."""
    data = read_law(code, target=target)
    write_workbook(data, out_path)
    return out_path, data
